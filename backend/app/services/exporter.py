

import io
import re

from app.models import Presentation
from app.services.llm.constants import SLIDE_HEIGHT, SLIDE_WIDTH


PDF_WIDTH = f"{SLIDE_WIDTH / 96:.3f}in"
PDF_HEIGHT = f"{SLIDE_HEIGHT / 96:.3f}in"


async def export_presentation(
        pres: Presentation,
        fmt: str = "pptx",
) -> bytes:

    fmt = (fmt or "pptx").lower()

    if fmt == "pdf":
        return await _export_pdf(pres)

    if fmt == "pptx":
        return await _export_pptx(pres)

    raise ValueError(
        "Unsupported export format. Use 'pdf' or 'pptx'."
    )


# ============================================================
# PDF
# ============================================================

async def _export_pdf(pres: Presentation) -> bytes:
    """
    Render HTML directly to PDF using Chromium.

    IMPORTANT:
    This does NOT screenshot the slides.

    Chromium's PDF engine preserves actual HTML text in the
    resulting PDF, so text can be selected and searched.
    """

    try:
        from playwright.async_api import async_playwright
    except ImportError as exc:
        raise RuntimeError(
            "Playwright is required for PDF export. "
            "Install with: pip install playwright && "
            "playwright install chromium"
        ) from exc

    async with async_playwright() as p:

        browser = await p.chromium.launch(
            headless=True,
            args=[
                "--no-sandbox",
                "--disable-setuid-sandbox",
            ],
        )

        page = await browser.new_page(
            viewport={
                "width": SLIDE_WIDTH,
                "height": SLIDE_HEIGHT,
            },
            device_scale_factor=1,
        )

        # Build one HTML document containing all slides.
        html = _build_pdf_document(pres)

        await page.set_content(
            html,
            wait_until="networkidle",
        )

        # Wait for fonts.
        await page.evaluate(
            """
            async () => {
                if (document.fonts) {
                    await document.fonts.ready;
                }
            }
            """
        )

        # Wait for images.
        await page.evaluate(
            """
            async () => {
                const images = Array.from(document.images);

                await Promise.all(
                    images.map(img => {
                        if (img.complete) {
                            return Promise.resolve();
                        }

                        return new Promise(resolve => {
                            img.onload = resolve;
                            img.onerror = resolve;
                        });
                    })
                );
            }
            """
        )

        await page.wait_for_timeout(100)

        pdf_bytes = await page.pdf(
            width=PDF_WIDTH,
            height=PDF_HEIGHT,

            # Very important:
            # preserve the CSS colors/backgrounds.
            print_background=True,

            # Remove browser margins.
            margin={
                "top": "0",
                "right": "0",
                "bottom": "0",
                "left": "0",
            },

            # Each .slide becomes one PDF page.
            prefer_css_page_size=True,
        )

        await browser.close()

        return pdf_bytes


def _build_pdf_document(pres: Presentation) -> str:
    """
    Build an HTML document specifically for PDF printing.

    Each generated slide.html remains HTML.
    We do NOT convert it into an image.
    """

    slides = []

    for index, slide in enumerate(pres.slides):

        slide_html = slide.html or ""

        if not slide_html.strip():
            slide_html = f"""
            <div class="empty-slide">
                {_escape_html(slide.label or "Slide")}
            </div>
            """

        slides.append(
            f"""
            <section
                class="slide"
                id="slide-{index + 1}"
            >
                {slide_html}
            </section>
            """
        )

    return f"""
<!DOCTYPE html>

<html lang="en">

<head>

<meta charset="UTF-8">

<title>{_escape_html(pres.title)}</title>

<style>

@page {{
    size: {PDF_WIDTH} {PDF_HEIGHT};
    margin: 0;
}}

html,
body {{
    margin: 0;
    padding: 0;
}}

body {{
    margin: 0;
    padding: 0;
}}

.slide {{
    width: {SLIDE_WIDTH}px;
    height: {SLIDE_HEIGHT}px;

    position: relative;

    overflow: hidden;

    page-break-after: always;

    break-after: page;

    box-sizing: border-box;
}}

.slide:last-child {{
    page-break-after: auto;
    break-after: auto;
}}

.empty-slide {{
    width: {SLIDE_WIDTH}px;
    height: {SLIDE_HEIGHT}px;

    display: flex;
    align-items: center;
    justify-content: center;

    font-family: Arial, sans-serif;
    font-size: 48px;
}}

* {{
    box-sizing: border-box;
}}

</style>

</head>

<body>

{''.join(slides)}

</body>

</html>
"""


# ============================================================
# PPTX
# ============================================================

async def _export_pptx(pres: Presentation) -> bytes:
    """
    PPTX export.

    Each HTML slide is rendered as a high-resolution image
    and inserted as a full-bleed PowerPoint slide.

    This gives maximum visual fidelity.
    """

    try:
        from playwright.async_api import async_playwright
    except ImportError as exc:
        raise RuntimeError(
            "Playwright is required for PPTX export."
        ) from exc

    try:
        from pptx import Presentation as PptxPresentation
        from pptx.util import Inches
    except ImportError as exc:
        raise RuntimeError(
            "python-pptx is required for PPTX export. "
            "Install with: pip install python-pptx"
        ) from exc

    async with async_playwright() as p:

        browser = await p.chromium.launch(
            headless=True,
            args=[
                "--no-sandbox",
                "--disable-setuid-sandbox",
            ],
        )

        page = await browser.new_page(
            viewport={
                "width": SLIDE_WIDTH,
                "height": SLIDE_HEIGHT,
            },
            device_scale_factor=2,
        )

        prs = PptxPresentation()

        # 16:9
        prs.slide_width = Inches(SLIDE_WIDTH / 96)
        prs.slide_height = Inches(SLIDE_HEIGHT / 96)

        blank_layout = prs.slide_layouts[6]

        for slide_data in pres.slides:

            slide_html = slide_data.html or ""

            if not slide_html.strip():
                slide_html = f"""
                <div style="
                    width:{SLIDE_WIDTH}px;
                    height:{SLIDE_HEIGHT}px;
                    display:flex;
                    align-items:center;
                    justify-content:center;
                    font-family:Arial;
                    font-size:48px;
                ">
                    {_escape_html(slide_data.label or "Slide")}
                </div>
                """

            document = _build_single_slide_document(
                slide_html
            )

            await page.set_content(
                document,
                wait_until="networkidle",
            )

            await page.evaluate(
                """
                async () => {
                    if (document.fonts) {
                        await document.fonts.ready;
                    }

                    const images = Array.from(document.images);

                    await Promise.all(
                        images.map(img => {
                            if (img.complete) {
                                return Promise.resolve();
                            }

                            return new Promise(resolve => {
                                img.onload = resolve;
                                img.onerror = resolve;
                            });
                        })
                    );
                }
                """
            )

            await page.wait_for_timeout(100)

            screenshot = await page.screenshot(
                type="png",
                full_page=False,
                animations="disabled",
            )

            pptx_slide = prs.slides.add_slide(
                blank_layout
            )

            # Full slide image.
            pptx_slide.shapes.add_picture(
                io.BytesIO(screenshot),
                0,
                0,
                width=prs.slide_width,
                height=prs.slide_height,
            )

            # Preserve speaker notes.
            if slide_data.notes:
                notes_slide = pptx_slide.notes_slide
                notes_slide.notes_text_frame.text = (
                    slide_data.notes
                )

        buffer = io.BytesIO()

        prs.save(buffer)

        await browser.close()

        return buffer.getvalue()


def _build_single_slide_document(
        slide_html: str,
) -> str:

    return f"""
<!DOCTYPE html>

<html>

<head>

<meta charset="UTF-8">

<style>

html,
body {{
    margin: 0;
    padding: 0;

    width: {SLIDE_WIDTH}px;
    height: {SLIDE_HEIGHT}px;

    overflow: hidden;
}}

body {{
    width: {SLIDE_WIDTH}px;
    height: {SLIDE_HEIGHT}px;
}}

* {{
    box-sizing: border-box;
}}

</style>

</head>

<body>

{slide_html}

</body>

</html>
"""


def _escape_html(value: str) -> str:
    return (
        (value or "")
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
        .replace("'", "&#39;")
    )
